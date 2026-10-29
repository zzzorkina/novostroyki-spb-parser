"""
Аналитика: сводные таблицы, графики, Excel-отчёт, карта и текст с выводами.
"""
from __future__ import annotations

import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # рисуем в файлы, без окна
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from cleaner import COLUMN_TITLES  # noqa: E402

# Единый стиль графиков: один спокойный цвет, приглушённая сетка
COLOR = "#2a78d6"
SURFACE = "#fcfcfb"
TEXT = "#0b0b0b"
TEXT_MUTED = "#52514e"
GRID = "#e6e5e1"
NO_DISTRICT = "Не указан"


# --- Сводные таблицы ---------------------------------------------------------------

def _share(part, total):
    return round(100 * part / total, 1) if total else None


def weighted_price(df: pd.DataFrame):
    """Средневзвешенная цена м² по жилой площади: большие объекты весят больше маленьких."""
    priced = df.dropna(subset=["price_m2", "living_area"])
    priced = priced[priced["living_area"] > 0]
    if priced.empty:
        return None
    return round((priced["price_m2"] * priced["living_area"]).sum() / priced["living_area"].sum())


def walkable_share(df: pd.DataFrame, max_minutes: float = 15):
    """Доля объектов, от которых до метро можно дойти пешком за max_minutes минут."""
    known = df.dropna(subset=["metro_minutes"])
    if known.empty or not df["metro_walk"].notna().any():
        return None
    walkable = known["metro_walk"].eq(True) & (known["metro_minutes"] <= max_minutes)
    return _share(walkable.sum(), len(known))


def is_residential(df: pd.DataFrame) -> pd.Series:
    """
    Жилой ли объект. Нежилые объекты (паркинги, коммерческие и социальные здания)
    не содержат квартир, поэтому в анализе рынка жилья их не учитываем.
    Если тип объекта сайт не указал, считаем объект жилым.
    """
    return ~(df["build_type"].notna() & (df["build_type"].str.lower() != "жилое"))


def non_residential_info(all_df: pd.DataFrame) -> dict:
    """Сводка по нежилым объектам — чтобы было видно, что и почему исключено."""
    other = all_df[~is_residential(all_df)]
    return {
        "Нежилых объектов (исключены из расчётов)": len(other),
        "ЖК, в которых есть нежилые объекты": other["complex"].nunique(),
        "Квартир в нежилых объектах": int(other["flats"].sum()),
        "Средняя этажность нежилых объектов": (round(other["floor_max"].mean(), 1)
                                              if other["floor_max"].notna().any() else None),
    }


def kpis(df: pd.DataFrame) -> dict:
    """Показатели рынка жилья. На вход подаются только жилые объекты."""
    flats = df["flats"].sum()
    area = df["living_area"].sum()
    return {
        "Жилых объектов строительства": len(df),
        "Жилых корпусов": int(df["buildings"].sum()) if df["buildings"].notna().any() else None,
        "Жилых комплексов": df["complex"].nunique(),
        "Групп компаний": df.loc[df["dev_group"] != "Не указан", "dev_group"].nunique(),
        "Застройщиков (юрлиц)": df.loc[df["developer"] != "Не указан", "developer"].nunique(),
        "Квартир": int(flats),
        "Жилая площадь, тыс. м²": round(area / 1000, 1),
        "Средняя площадь квартиры, м²": round(area / flats, 1) if flats else None,
        "Средняя этажность (макс.)": round(df["floor_max"].mean(), 1) if df["floor_max"].notna().any() else None,
        "Доля объектов с прошедшим плановым сроком, %": _share(df["deadline_passed"].eq(True).sum(), len(df)),
        "Средневзвешенная цена м², ₽": weighted_price(df),
        "Доля объектов с указанной ценой, %": _share(df["price_m2"].notna().sum(), len(df)),
        "Среднее время до метро, мин": (round(df["metro_minutes"].mean(), 1)
                                        if df["metro_minutes"].notna().any() else None),
        "Доля объектов в 15 минутах пешком от метро, %": walkable_share(df),
    }


def district_coverage(df: pd.DataFrame) -> float:
    """Доля объектов, у которых удалось определить район."""
    return float((df["district"] != NO_DISTRICT).mean()) if len(df) else 0.0


def group_summary(df: pd.DataFrame, by: str) -> pd.DataFrame:
    """Сводка по одному измерению: объекты, корпуса, ЖК, квартиры, площадь, доля рынка, цена."""
    d = df.copy()
    has_price = d["price_m2"].notna() & d["living_area"].gt(0)
    d["_price_x_area"] = (d["price_m2"] * d["living_area"]).where(has_price)
    d["_priced_area"] = d["living_area"].where(has_price)
    g = d.groupby(by, dropna=True)
    table = pd.DataFrame({
        "Объектов": g.size(),
        "Жилых корпусов": g["buildings"].sum(min_count=1),
        "ЖК": g["complex"].nunique(),
        "Квартир": g["flats"].sum(),
        "Жилая площадь, м²": g["living_area"].sum().round(0),
        "Средняя этажность (макс.)": g["floor_max"].mean().round(1),
    })
    table["Средняя площадь квартиры, м²"] = (table["Жилая площадь, м²"] / table["Квартир"]).where(
        table["Квартир"] > 0).round(1)
    table["Доля квартир, %"] = (100 * table["Квартир"] / table["Квартир"].sum()).round(1)
    if has_price.any():
        table["Средневзвешенная цена м², ₽"] = (g["_price_x_area"].sum(min_count=1)
                                               / g["_priced_area"].sum(min_count=1)).round(0)
    return table.sort_values("Квартир", ascending=False)


def build_tables(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    tables = {
        "По группам компаний": group_summary(df, "dev_group"),
        "По застройщикам (юрлица)": group_summary(df, "developer"),
        "По годам сдачи": group_summary(df, "ready_year").sort_index(),
    }
    if district_coverage(df) >= 0.3:   # если районов в адресах почти нет, таблица бессмысленна
        tables["По районам"] = group_summary(df, "district")
    if df["metro"].notna().any():
        tables["По метро"] = group_summary(df, "metro")
    if df["housing_class"].notna().any():
        tables["По классу жилья"] = group_summary(df, "housing_class")
    index_names = {"По группам компаний": "Группа компаний", "По застройщикам (юрлица)": "Застройщик",
                   "По годам сдачи": "Год сдачи", "По районам": "Район", "По метро": "Ближайшее метро",
                   "По типу объекта": "Тип объекта", "По классу жилья": "Класс"}
    for name in tables:
        tables[name].index.name = index_names[name]
    return tables


# --- Графики ------------------------------------------------------------------------

def _style(ax, title: str, subtitle: str = ""):
    ax.set_facecolor(SURFACE)
    ax.figure.set_facecolor(SURFACE)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color(GRID)
    ax.tick_params(colors=TEXT_MUTED, labelsize=10, length=0)
    ax.figure.text(0.01, 0.97, title, fontsize=14, weight="bold", color=TEXT, va="top")
    if subtitle:
        ax.figure.text(0.01, 0.915, subtitle, fontsize=10, color=TEXT_MUTED, va="top")


def _fmt(n: float) -> str:
    return f"{n:,.0f}".replace(",", " ")


def hbar(series: pd.Series, title: str, subtitle: str, path: Path, value_fmt=_fmt):
    """Горизонтальные столбики: категории сверху вниз от большего к меньшему."""
    series = series.dropna().sort_values()
    if series.empty:
        return None
    height = max(3.2, 0.42 * len(series) + 1.4)
    fig, ax = plt.subplots(figsize=(9, height))
    bars = ax.barh(series.index.astype(str), series.values, color=COLOR, height=0.62)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: _fmt(x)))
    vmax = series.max()
    for bar, value in zip(bars, series.values):
        ax.text(bar.get_width() + vmax * 0.01, bar.get_y() + bar.get_height() / 2,
                value_fmt(value), va="center", fontsize=9, color=TEXT_MUTED)
    ax.set_xlim(0, vmax * 1.12)
    _style(ax, title, subtitle)
    fig.tight_layout(rect=(0, 0, 1, 0.88))
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def vbar(series: pd.Series, title: str, subtitle: str, path: Path):
    series = series.dropna()
    if series.empty:
        return None
    fig, ax = plt.subplots(figsize=(9, 4.6))
    labels = [str(i) for i in series.index]
    bars = ax.bar(labels, series.values, color=COLOR, width=0.6)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: _fmt(x)))
    vmax = series.max()
    for bar, value in zip(bars, series.values):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + vmax * 0.015,
                _fmt(value), ha="center", fontsize=9, color=TEXT_MUTED)
    ax.set_ylim(0, vmax * 1.12)
    _style(ax, title, subtitle)
    fig.tight_layout(rect=(0, 0, 1, 0.86))
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def make_charts(df: pd.DataFrame, tables: dict, out_dir: Path, region_name: str) -> list[Path]:
    """Строит графики. Если для графика нет данных, он просто пропускается."""
    out_dir.mkdir(parents=True, exist_ok=True)
    charts = []

    top_groups = tables["По группам компаний"]["Квартир"].drop("Не указан", errors="ignore").head(10)
    charts.append(hbar(top_groups, "Топ-10 групп компаний по числу строящихся квартир",
                       f"{region_name}, квартир в строящихся объектах", out_dir / "01_top_groups.png"))

    charts.append(vbar(tables["По годам сдачи"]["Квартир"], "Когда сдадут квартиры",
                       f"{region_name}, квартир по плановому году сдачи", out_dir / "02_flats_by_year.png"))

    # Районы показываем, только если их удалось определить хотя бы у трети объектов
    if district_coverage(df) >= 0.3:
        by_district = tables["По районам"]["Квартир"].drop(NO_DISTRICT, errors="ignore")
        charts.append(hbar(by_district, "Где строят больше всего квартир",
                           f"{region_name}, квартир по районам", out_dir / "04_flats_by_district.png"))
        area = tables["По районам"]
        area = area[area["Объектов"] >= 3].drop(NO_DISTRICT, errors="ignore")["Средняя площадь квартиры, м²"]
        charts.append(hbar(area, "Средняя площадь квартиры по районам",
                           "м², жилая площадь / число квартир; районы с 3+ объектами",
                           out_dir / "05_avg_flat_area_by_district.png",
                           value_fmt=lambda v: f"{v:.1f}".replace(".", ",")))

    if "По метро" in tables:
        charts.append(hbar(tables["По метро"]["Квартир"].head(15), "Топ-15 станций метро по числу строящихся квартир",
                           f"{region_name}, квартир в объектах, для которых станция — ближайшая",
                           out_dir / "03_flats_by_metro.png"))
    return [c for c in charts if c]


# --- Карта ----------------------------------------------------------------------------

def make_map(df: pd.DataFrame, path: Path) -> Path | None:
    """Интерактивная карта: точка — дом, размер — число квартир."""
    try:
        import folium
    except ImportError:
        print("  (folium не установлен — карту пропускаю)")
        return None
    points = df.dropna(subset=["lat", "lon"])
    if points.empty:
        print("  (в данных нет координат — карту пропускаю)")
        return None
    fmap = folium.Map(location=[points["lat"].mean(), points["lon"].mean()],
                      zoom_start=10, tiles="OpenStreetMap")
    for _, row in points.iterrows():
        flats = row["flats"] if pd.notna(row["flats"]) else 0
        popup = (f"<b>{row['complex'] or 'ЖК без названия'}</b><br>"
                 f"{row['dev_group'] or ''} ({row['developer'] or ''})<br>{row['address'] or ''}<br>"
                 f"Квартир: {_fmt(flats)}<br>Срок сдачи: "
                 f"{row['ready_date'].strftime('%d.%m.%Y') if pd.notna(row['ready_date']) else '—'}")
        folium.CircleMarker(
            location=[row["lat"], row["lon"]],
            radius=3 + math.sqrt(flats) / 4,
            color="#ffffff", weight=1, fill=True, fill_color=COLOR, fill_opacity=0.75,
            popup=folium.Popup(popup, max_width=300),
        ).add_to(fmap)
    fmap.save(str(path))
    return path


# --- Excel ------------------------------------------------------------------------------

def _format_sheet(ws):
    header_fill = PatternFill("solid", fgColor="EAF2FC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for idx, column in enumerate(ws.columns, start=1):
        values = [str(c.value) for c in column[:200] if c.value is not None]
        width = min(max([len(v) for v in values] + [8]) + 2, 50)
        ws.column_dimensions[get_column_letter(idx)].width = width


def for_export(df: pd.DataFrame) -> pd.DataFrame:
    """Копия таблицы для выгрузки: даты без времени, русские названия колонок."""
    out = df.copy()
    out["ready_date"] = out["ready_date"].dt.date
    return out.rename(columns=COLUMN_TITLES)


def write_excel(res: pd.DataFrame, all_df: pd.DataFrame, tables: dict, info: dict, path: Path) -> Path:
    """res — жилые объекты (для показателей), all_df — все объекты (лист «Объекты»)."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        rows = list(info.items()) + list(kpis(res).items()) + list(non_residential_info(all_df).items())
        summary = pd.DataFrame(rows, columns=["Показатель", "Значение"])
        summary.to_excel(writer, sheet_name="Сводка", index=False)
        for_export(all_df).to_excel(writer, sheet_name="Объекты", index=False)
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name)
        fields = pd.DataFrame(list(COLUMN_TITLES.items()), columns=["Поле в коде", "Название в отчёте"])
        fields.to_excel(writer, sheet_name="Описание полей", index=False)

        for ws in writer.book.worksheets:
            _format_sheet(ws)
            ws.freeze_panes = "A2" if ws.title in ("Сводка", "Описание полей") else "B2"
    return path


# --- Текстовые выводы ------------------------------------------------------------------

def _dec(x, digits: int = 1) -> str:
    """52.04 -> '52,0' — десятичная запятая, как принято в русском тексте."""
    return f"{x:,.{digits}f}".replace(",", " ").replace(".", ",")


def _in_district(name: str) -> str:
    """'Приморский' -> 'в Приморском районе'."""
    if name.endswith(("ий", "ый")):
        return f"в {name[:-2]}ом районе"
    return f"в районе {name}"


def write_summary(df: pd.DataFrame, all_df: pd.DataFrame, tables: dict, info: dict, path: Path) -> Path:
    """df — жилые объекты, all_df — все объекты каталога."""
    k = kpis(df)
    nr = non_residential_info(all_df)
    lines = [f"# Новостройки: {info['Регион']} — ключевые выводы", ""]
    if info.get("Режим") == "ДЕМО":
        lines += ["> ⚠️ Отчёт построен на **демо-данных** (выдуманные объекты для проверки кода). "
                  "Запустите `python main.py`, чтобы собрать реальные данные.", ""]
    lines += [f"Источник: {info['Источник']}. Данные на {info['Дата сбора']}.", ""]

    buildings = f" (жилых корпусов — {_fmt(k['Жилых корпусов'])})" if k["Жилых корпусов"] else ""
    lines.append(f"- **Объём рынка.** Жилых объектов строительства — {k['Жилых объектов строительства']}{buildings}, "
                 f"жилых комплексов — {k['Жилых комплексов']}. В них {_fmt(k['Квартир'])} квартир "
                 f"и {_dec(k['Жилая площадь, тыс. м²'])} тыс. м² жилой площади.")
    if nr["Нежилых объектов (исключены из расчётов)"]:
        floors = nr["Средняя этажность нежилых объектов"]
        flats_text = ("квартир и жилой площади в них нет" if not nr["Квартир в нежилых объектах"]
                      else f"квартир в них {_fmt(nr['Квартир в нежилых объектах'])}")
        guess = (f" Средняя этажность — {_dec(floors)}, так что это, вероятно, паркинги, коммерческие "
                 "и социальные здания." if floors and floors <= 10 else "")
        lines.append(f"- **Нежилые объекты.** Ещё {nr['Нежилых объектов (исключены из расчётов)']} объектов каталога "
                     f"нежилые (в {nr['ЖК, в которых есть нежилые объекты']} ЖК): {flats_text}.{guess} "
                     "В остальных расчётах они не участвуют.")

    groups = tables["По группам компаний"].drop("Не указан", errors="ignore")
    if len(groups):
        top5 = groups.head(5)
        names = ", ".join(f"{n} — {_dec(s)}%" for n, s in zip(top5.index, top5["Доля квартир, %"]))
        lines.append(f"- **Участники.** На рынке {k['Групп компаний']} групп компаний и "
                     f"{k['Застройщиков (юрлиц)']} юрлиц-застройщиков: крупные девелоперы оформляют проекты "
                     "на отдельные специализированные компании (СЗ), поэтому анализ ведётся по группам.")
        lines.append(f"- **Концентрация.** На топ-5 групп приходится {_dec(top5['Доля квартир, %'].sum())}% "
                     f"строящихся квартир: {names}.")

    dist = tables.get("По районам", pd.DataFrame()).drop(NO_DISTRICT, errors="ignore")
    if "По метро" in tables and district_coverage(df) < 0.5:
        top3 = tables["По метро"].head(3)
        names = ", ".join(f"{n} — {_dec(s)}%" for n, s in zip(top3.index, top3["Доля квартир, %"]))
        lines.append(f"- **География.** Больше всего квартир строится рядом со станциями метро: {names}.")
    elif len(dist):
        top3 = dist.head(3)
        names = ", ".join(f"{n} — {_dec(s)}%" for n, s in zip(top3.index, top3["Доля квартир, %"]))
        lines.append(f"- **География.** Больше всего квартир строится в районах: {names}.")
        big = dist[dist["Объектов"] >= 3]
        if len(big) >= 2 and k["Средняя площадь квартиры, м²"]:
            a = big["Средняя площадь квартиры, м²"]
            lines.append(f"- **Формат квартир.** Средняя площадь квартиры — {_dec(k['Средняя площадь квартиры, м²'])} м². "
                         f"Самые просторные квартиры {_in_district(a.idxmax())} ({_dec(a.max())} м²), "
                         f"самые компактные — {_in_district(a.idxmin())} ({_dec(a.min())} м²).")

    if k["Среднее время до метро, мин"]:
        walk = k["Доля объектов в 15 минутах пешком от метро, %"]
        walk_text = f" В 15 минутах пешком от метро — {_dec(walk)}% объектов." if walk is not None else ""
        lines.append(f"- **Транспорт.** Среднее время до ближайшего метро — "
                     f"{_dec(k['Среднее время до метро, мин'])} мин.{walk_text}")

    years = tables["По годам сдачи"]
    if len(years):
        peak = years["Квартир"].idxmax()
        lines.append(f"- **Сроки.** Пик ввода — {peak} год: {_fmt(years.loc[peak, 'Квартир'])} квартир, "
                     f"или {_dec(years.loc[peak, 'Доля квартир, %'])}% всех строящихся.")
    if k["Доля объектов с прошедшим плановым сроком, %"]:
        lines.append(f"- **Риски.** У {_dec(k['Доля объектов с прошедшим плановым сроком, %'])}% объектов плановый срок "
                     "сдачи уже прошёл, а они всё ещё числятся строящимися — кандидаты на проверку переноса сроков.")
    if k["Средневзвешенная цена м², ₽"]:
        lines.append(f"- **Цены.** Средневзвешенная (по жилой площади) цена м² — "
                     f"{_fmt(k['Средневзвешенная цена м², ₽'])} ₽; цена указана у "
                     f"{_dec(k['Доля объектов с указанной ценой, %'])}% объектов.")

    unknown = (df["district"] == NO_DISTRICT).mean()
    district_note = (f"- Район не удалось определить по адресу у {unknown:.0%} объектов."
                     if unknown < 0.5 or not df["metro"].notna().any() else
                     "- В адресах на сайте район обычно не указан, поэтому география показана по ближайшему метро.")
    lines += ["", "## Качество данных и методика", district_note,
              f"- Нет числа квартир у {df['flats'].isna().mean():.0%} объектов, "
              f"нет координат у {df['lat'].isna().mean():.0%}.",
              "- Объект строительства — единица учёта ЕИСЖС; в одном объекте может быть несколько корпусов. "
              "ЖК считаются по уникальным коммерческим названиям.",
              "- Показатели рынка жилья считаются только по жилым объектам; нежилые остаются в Excel "
              "на листах «Объекты» и «По типу объекта».",
              "- Если группа компаний на сайте не указана, застройщик считается самостоятельной группой."]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def run(df: pd.DataFrame, info: dict, out_dir: Path) -> dict:
    """Строит все выходные файлы и возвращает пути к ним."""
    out_dir.mkdir(parents=True, exist_ok=True)
    res = df[is_residential(df)]           # рынок жилья — только жилые объекты
    tables = build_tables(res)
    if df["build_type"].notna().any():
        tables["По типу объекта"] = group_summary(df, "build_type")
    region = info["Регион"]
    result = {
        "excel": write_excel(res, df, tables, info, out_dir / "novostroyki_report.xlsx"),
        "csv": out_dir / "objects.csv",
        "charts": make_charts(res, tables, out_dir / "charts", region),
        "map": make_map(res, out_dir / "map.html"),
        "summary": write_summary(res, df, tables, info, out_dir / "summary.md"),
    }
    # CSV для Power BI / Google Sheets: UTF-8 с BOM, чтобы Excel не ломал кириллицу
    for_export(df).to_csv(result["csv"], index=False, encoding="utf-8-sig")
    return result

